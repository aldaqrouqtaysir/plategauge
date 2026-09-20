"""LeFood-Set manifest construction, hashing, duplicate checks, and data gate audit."""

from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

import numpy as np
from PIL import Image, ImageOps

from .errors import DataIntegrityError, missing_extra
from .folds import (
    EXPECTED_FOLD_COUNTS,
    fold_for_category,
    validate_duplicate_groups,
    validate_observed_counts,
)
from .schema import ManifestRecord, read_manifest, write_manifest

DEFAULT_DATASET_ROOT = Path("data/raw/lefood-v1/LeFood-Set Leftovers Food Dataset/LeFood-Set")
DEFAULT_MANIFEST = Path("data/manifests/lefood_v1_manifest.csv")
DEFAULT_SOURCE_AUDIT = Path("data/manifests/lefood_v1_source_audit.csv")
DEFAULT_AUDIT_REPORT = Path("data/manifests/lefood_v1_audit.json")
DEFAULT_DUPLICATE_REVIEW = Path("data/manifests/lefood_v1_duplicate_candidates.csv")

SOURCE_AUDIT_FIELDS = (
    "source_row",
    "dataset_id",
    "food_name",
    "category",
    "before_filename",
    "before_mass_g",
    "after_filename",
    "after_mass_g",
    "observer_score",
    "status",
    "reason",
)


@dataclass(frozen=True, slots=True)
class AuditReport:
    """Machine-readable result of the preregistered data gate."""

    schema_version: str
    manifest_sha256: str | None
    workbook_rows: int
    matched_pairs: int
    missing_image_rows: int
    valid_pairs: int
    invalid_mass_pairs: int
    valid_categories: int
    fold_counts: dict[str, int]
    exact_duplicate_components: int
    near_duplicate_edges: int
    duplicate_components: int
    cross_fold_duplicate_components: int
    files_verified: bool
    hashes_verified: bool
    expected_counts_match: bool
    passed: bool
    issues: tuple[str, ...]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def sha256_file(path: str | Path, *, chunk_size: int = 1024 * 1024) -> str:
    """Hash a file without loading it entirely into memory."""

    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _open_rgb(path: Path) -> Image.Image:
    with Image.open(path) as image:
        return ImageOps.exif_transpose(image).convert("RGB")


def _image_metadata(path: Path) -> tuple[int, int, str]:
    with Image.open(path) as image:
        image.verify()
    with Image.open(path) as image:
        width, height = ImageOps.exif_transpose(image).size
    return width, height, sha256_file(path)


def _phash_and_thumb(path: Path) -> tuple[int, np.ndarray]:
    """Return a 64-bit perceptual hash and a normalized 64px grayscale thumbnail."""

    image = _open_rgb(path).convert("L")
    small = np.asarray(image.resize((32, 32), Image.Resampling.LANCZOS), dtype=np.float64)
    n = 32
    positions = np.arange(n, dtype=np.float64)
    frequencies = positions[:, None]
    basis = np.cos(np.pi * (2.0 * positions + 1.0) * frequencies / (2.0 * n))
    basis[0] *= math.sqrt(1.0 / n)
    basis[1:] *= math.sqrt(2.0 / n)
    low = (basis @ small @ basis.T)[:8, :8].reshape(-1)
    median = float(np.median(low[1:]))
    value = 0
    for index, bit in enumerate(low >= median):
        if bool(bit):
            value |= 1 << index
    thumb = np.asarray(image.resize((64, 64), Image.Resampling.LANCZOS), dtype=np.float64) / 255.0
    return value, thumb


def _windowed_ssim(left: np.ndarray, right: np.ndarray) -> float:
    """Compute mean local SSIM with the standard 11px Gaussian window.

    A global moment approximation is unsafe for duplicate detection: two mostly
    uniform plate photographs can have very similar aggregate moments despite
    different local content.
    """

    if left.shape != right.shape or left.ndim != 2:
        raise DataIntegrityError("SSIM inputs must be identically shaped grayscale arrays")
    coordinates = np.arange(11, dtype=np.float64) - 5.0
    kernel_1d = np.exp(-(coordinates**2) / (2.0 * 1.5**2))
    kernel_1d /= kernel_1d.sum()
    kernel = np.outer(kernel_1d, kernel_1d)

    def filter_image(image: np.ndarray) -> np.ndarray:
        padded = np.pad(image, 5, mode="reflect")
        windows = np.lib.stride_tricks.sliding_window_view(padded, (11, 11))
        filtered: np.ndarray = np.einsum("ijkl,kl->ij", windows, kernel, optimize=True)
        return filtered

    mean_left = filter_image(left)
    mean_right = filter_image(right)
    variance_left = np.maximum(0.0, filter_image(left * left) - mean_left * mean_left)
    variance_right = np.maximum(0.0, filter_image(right * right) - mean_right * mean_right)
    covariance = filter_image(left * right) - mean_left * mean_right
    c1, c2 = 0.01**2, 0.03**2
    numerator = (2.0 * mean_left * mean_right + c1) * (2.0 * covariance + c2)
    denominator = (mean_left**2 + mean_right**2 + c1) * (variance_left + variance_right + c2)
    return float(
        np.mean(
            np.divide(numerator, denominator, out=np.ones_like(numerator), where=denominator != 0)
        )
    )


class _UnionFind:
    def __init__(self, items: list[str]) -> None:
        self.parent = {item: item for item in items}

    def find(self, item: str) -> str:
        parent = self.parent[item]
        if parent != item:
            self.parent[item] = self.find(parent)
        return self.parent[item]

    def union(self, left: str, right: str) -> None:
        root_left, root_right = self.find(left), self.find(right)
        if root_left != root_right:
            self.parent[max(root_left, root_right)] = min(root_left, root_right)


def find_duplicate_groups(
    image_paths: list[Path],
    *,
    phash_distance: int = 4,
    ssim_threshold: float = 0.995,
) -> tuple[dict[str, str], int, int, list[dict[str, str]]]:
    """Find exact and confirmed near-duplicate image components.

    The quadratic pHash comparison is deliberate and tractable for LeFood's 1,048
    images. SSIM is computed only for pHash candidates.
    """

    keys = [path.as_posix() for path in image_paths]
    union_find = _UnionFind(keys)
    exact_buckets: dict[str, list[str]] = defaultdict(list)
    phashes: dict[str, int] = {}
    thumbs: dict[str, np.ndarray] = {}
    for path, key in zip(image_paths, keys, strict=True):
        exact_buckets[sha256_file(path)].append(key)
        phashes[key], thumbs[key] = _phash_and_thumb(path)

    exact_components = 0
    candidates: list[dict[str, str]] = []
    for members in exact_buckets.values():
        if len(members) > 1:
            exact_components += 1
            for member in members[1:]:
                union_find.union(members[0], member)
                candidates.append(
                    {
                        "left_path": members[0],
                        "right_path": member,
                        "match_type": "exact_sha256",
                        "phash_distance": "0",
                        "ssim": "1",
                        "duplicate_group": "",
                    }
                )

    near_edges = 0
    for left_index, left in enumerate(keys):
        for right in keys[left_index + 1 :]:
            if phashes[left] ^ phashes[right] == 0 and union_find.find(left) == union_find.find(
                right
            ):
                continue
            if (phashes[left] ^ phashes[right]).bit_count() <= phash_distance:
                ssim = _windowed_ssim(thumbs[left], thumbs[right])
                if ssim >= ssim_threshold:
                    if union_find.find(left) != union_find.find(right):
                        near_edges += 1
                    union_find.union(left, right)
                    candidates.append(
                        {
                            "left_path": left,
                            "right_path": right,
                            "match_type": "phash_ssim",
                            "phash_distance": str((phashes[left] ^ phashes[right]).bit_count()),
                            "ssim": format(ssim, ".10f"),
                            "duplicate_group": "",
                        }
                    )

    components: dict[str, list[str]] = defaultdict(list)
    for key in keys:
        components[union_find.find(key)].append(key)
    duplicate_components = [sorted(members) for members in components.values() if len(members) > 1]
    assignments: dict[str, str] = {}
    for members in sorted(duplicate_components):
        group_hash = hashlib.sha256("\n".join(members).encode()).hexdigest()[:12]
        for member in members:
            assignments[member] = f"dup-{group_hash}"
    for candidate in candidates:
        candidate["duplicate_group"] = assignments[candidate["left_path"]]
    return assignments, exact_components, near_edges, candidates


def _load_workbook_rows(workbook_path: Path) -> list[tuple[Any, ...]]:
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover - environment-dependent
        raise missing_extra("Excel dataset ingestion", "data", "openpyxl") from exc

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.active
    values = list(sheet.values)
    workbook.close()
    expected_header = (
        "ID",
        "Name of the food",
        "Image Before Eaten",
        "Weight Before Eaten (g)",
        "Image After Eaten",
        "Weight After Eaten (g)",
        "Visual Estimation by Observer (1-7)",
    )
    if not values or tuple(values[0]) != expected_header:
        raise DataIntegrityError(
            "LeFood workbook header changed; refusing heuristic column matching"
        )
    return [tuple(row) for row in values[1:]]


def build_lefood_manifest(
    dataset_root: str | Path = DEFAULT_DATASET_ROOT,
    *,
    detect_near_duplicates: bool = True,
) -> tuple[list[ManifestRecord], list[dict[str, str]], dict[str, Any]]:
    """Build matched-pair records directly from the official v1 workbook and archive."""

    root = Path(dataset_root)
    workbook_path = root / "data_original.xlsx"
    image_root = root / "leftover dataset"
    if not workbook_path.is_file() or not image_root.is_dir():
        raise DataIntegrityError(
            f"Expected data_original.xlsx and 'leftover dataset' under {root.resolve()}"
        )
    workbook_rows = _load_workbook_rows(workbook_path)
    if len(workbook_rows) != 678:
        raise DataIntegrityError(f"Expected 678 workbook rows, found {len(workbook_rows)}")

    staged: list[dict[str, Any]] = []
    source_audit: list[dict[str, str]] = []
    all_image_paths: list[Path] = []
    metadata_cache: dict[Path, tuple[int, int, str]] = {}
    for offset, row in enumerate(workbook_rows, start=2):
        dataset_id, food_name, before_name, before_mass, after_name, after_mass, observer = row
        before_filename = str(before_name)
        after_filename = str(after_name)
        category = before_filename.split("_", 1)[0] if before_filename else ""
        before_path = image_root / "data_before" / category / before_filename
        after_path = image_root / "data_after" / category / after_filename
        missing = [path.name for path in (before_path, after_path) if not path.is_file()]
        source_row = {
            "source_row": str(offset),
            "dataset_id": str(dataset_id),
            "food_name": str(food_name),
            "category": category,
            "before_filename": before_filename,
            "before_mass_g": str(before_mass),
            "after_filename": after_filename,
            "after_mass_g": str(after_mass),
            "observer_score": str(observer),
            "status": "",
            "reason": "",
        }
        if missing:
            source_row["status"] = "missing_image"
            source_row["reason"] = "missing:" + ",".join(missing)
            source_audit.append(source_row)
            continue
        if after_filename.split("_", 1)[0] != category:
            raise DataIntegrityError(f"Category mismatch in workbook row {offset}")
        before_mass_float = float(before_mass)
        after_mass_float = float(after_mass)
        if before_mass_float <= 0 or after_mass_float < 0:
            raise DataIntegrityError(f"Impossible mass in workbook row {offset}")
        is_valid = after_mass_float <= before_mass_float
        source_row["status"] = "matched_valid" if is_valid else "matched_excluded"
        source_row["reason"] = "" if is_valid else "after_mass_exceeds_before_mass"
        source_audit.append(source_row)
        for path in (before_path, after_path):
            if path not in metadata_cache:
                metadata_cache[path] = _image_metadata(path)
                all_image_paths.append(path)
        staged.append(
            {
                "dataset_id": int(dataset_id),
                "source_row": offset,
                "food_name": str(food_name),
                "category": category,
                "before_path": before_path,
                "after_path": after_path,
                "before_mass_g": before_mass_float,
                "after_mass_g": after_mass_float,
                "observer_score": int(observer),
                "is_valid": is_valid,
            }
        )

    duplicate_assignments: dict[str, str] = {}
    exact_components = 0
    near_components = 0
    duplicate_candidates: list[dict[str, str]] = []
    if detect_near_duplicates:
        (
            duplicate_assignments,
            exact_components,
            near_components,
            duplicate_candidates,
        ) = find_duplicate_groups(all_image_paths)

    records: list[ManifestRecord] = []
    root_resolved = root.resolve()
    for item in staged:
        before_path = item["before_path"]
        after_path = item["after_path"]
        before_width, before_height, before_hash = metadata_cache[before_path]
        after_width, after_height, after_hash = metadata_cache[after_path]
        duplicate_groups = sorted(
            {
                duplicate_assignments.get(before_path.as_posix(), ""),
                duplicate_assignments.get(after_path.as_posix(), ""),
            }
            - {""}
        )
        duplicate_group = "+".join(duplicate_groups)
        category = item["category"]
        records.append(
            ManifestRecord(
                sample_id=f"lefood-{item['dataset_id']:04d}",
                source_row=item["source_row"],
                food_name=item["food_name"],
                category=category,
                before_path=before_path.resolve().relative_to(root_resolved).as_posix(),
                after_path=after_path.resolve().relative_to(root_resolved).as_posix(),
                before_mass_g=item["before_mass_g"],
                after_mass_g=item["after_mass_g"],
                leftover_fraction=item["after_mass_g"] / item["before_mass_g"],
                observer_score=item["observer_score"],
                before_width=before_width,
                before_height=before_height,
                after_width=after_width,
                after_height=after_height,
                before_sha256=before_hash,
                after_sha256=after_hash,
                duplicate_group=duplicate_group,
                is_valid=item["is_valid"],
                exclusion_reason="" if item["is_valid"] else "after_mass_exceeds_before_mass",
                outer_fold=fold_for_category(category),
            )
        )
    diagnostics = {
        "exact_duplicate_components": exact_components,
        "near_duplicate_edges": near_components,
        "duplicate_candidates": duplicate_candidates,
    }
    return records, source_audit, diagnostics


def write_source_audit(rows: list[dict[str, str]], path: str | Path) -> None:
    audit_path = Path(path)
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    with audit_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=SOURCE_AUDIT_FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def write_duplicate_review(
    rows: list[dict[str, str]], path: str | Path, *, dataset_root: str | Path
) -> None:
    """Write every threshold-passing pair for human confirmation."""

    review_path = Path(path)
    review_path.parent.mkdir(parents=True, exist_ok=True)
    root = Path(dataset_root).resolve()
    fields = (
        "left_path",
        "right_path",
        "match_type",
        "phash_distance",
        "ssim",
        "duplicate_group",
    )
    normalized: list[dict[str, str]] = []
    for row in rows:
        converted = dict(row)
        for column in ("left_path", "right_path"):
            candidate = Path(converted[column]).resolve()
            converted[column] = candidate.relative_to(root).as_posix()
        normalized.append(converted)
    with review_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(
            sorted(normalized, key=lambda row: (row["duplicate_group"], row["left_path"]))
        )


def audit_manifest(
    records: list[ManifestRecord],
    *,
    dataset_root: str | Path | None = None,
    workbook_rows: int = 678,
    missing_image_rows: int = 154,
    exact_duplicate_components: int = 0,
    near_duplicate_components: int = 0,
    verify_hashes: bool = False,
    manifest_sha256: str | None = None,
) -> AuditReport:
    """Run all count, fold, duplicate, file, and optional hash gates."""

    issues: list[str] = []
    valid = [record for record in records if record.is_valid]
    invalid = [record for record in records if not record.is_valid]
    crossing_groups: dict[str, set[int]] = defaultdict(set)
    for record in valid:
        if record.duplicate_group:
            for group in record.duplicate_group.split("+"):
                crossing_groups[group].add(record.outer_fold)
    expected_counts_match = True
    try:
        validate_observed_counts((record.category for record in valid), exact=True)
    except DataIntegrityError as exc:
        expected_counts_match = False
        issues.append(str(exc))
    try:
        validate_duplicate_groups(
            (record.category for record in valid),
            (record.duplicate_group for record in valid),
            (record.outer_fold for record in valid),
        )
    except DataIntegrityError as exc:
        issues.append(str(exc))

    root = Path(dataset_root) if dataset_root is not None else None
    files_verified = root is not None
    hashes_verified = root is not None and verify_hashes
    if root is None:
        issues.append("Dataset files were not verified against the manifest")
    elif not verify_hashes:
        issues.append("Dataset image hashes were not re-verified")
    if root is not None:
        for record in records:
            for relative_path, expected_hash in (
                (record.before_path, record.before_sha256),
                (record.after_path, record.after_sha256),
            ):
                path = root / relative_path
                if not path.is_file():
                    issues.append(f"Missing manifest image: {relative_path}")
                    continue
                if verify_hashes and sha256_file(path) != expected_hash:
                    issues.append(f"Hash mismatch: {relative_path}")

    actual = {
        "workbook_rows": workbook_rows,
        "matched_pairs": len(records),
        "missing_image_rows": missing_image_rows,
        "valid_pairs": len(valid),
        "invalid_mass_pairs": len(invalid),
        "valid_categories": len({record.category for record in valid}),
    }
    expected = {
        "workbook_rows": 678,
        "matched_pairs": 524,
        "missing_image_rows": 154,
        "valid_pairs": 514,
        "invalid_mass_pairs": 10,
        "valid_categories": 34,
    }
    for name, expected_value in expected.items():
        if actual[name] != expected_value:
            issues.append(f"{name}: expected {expected_value}, found {actual[name]}")
    invalid_reasons = Counter(record.exclusion_reason for record in invalid)
    if invalid_reasons and invalid_reasons != Counter({"after_mass_exceeds_before_mass": 10}):
        issues.append(f"Unexpected exclusions: {dict(invalid_reasons)!r}")
    fold_counts = Counter(record.outer_fold for record in valid)
    if dict(sorted(fold_counts.items())) != EXPECTED_FOLD_COUNTS:
        issues.append(f"Fold counts changed: {dict(sorted(fold_counts.items()))!r}")
    if len(valid) < 500:
        issues.append("Data stop rule triggered: fewer than 500 valid pairs")

    return AuditReport(
        schema_version="1.0",
        manifest_sha256=manifest_sha256,
        workbook_rows=workbook_rows,
        matched_pairs=len(records),
        missing_image_rows=missing_image_rows,
        valid_pairs=len(valid),
        invalid_mass_pairs=len(invalid),
        valid_categories=len({record.category for record in valid}),
        fold_counts={str(key): value for key, value in sorted(fold_counts.items())},
        exact_duplicate_components=exact_duplicate_components,
        near_duplicate_edges=near_duplicate_components,
        duplicate_components=len(crossing_groups),
        cross_fold_duplicate_components=sum(
            len(group_folds) > 1 for group_folds in crossing_groups.values()
        ),
        files_verified=files_verified,
        hashes_verified=hashes_verified,
        expected_counts_match=expected_counts_match,
        passed=not issues,
        issues=tuple(issues),
    )


def create_and_audit_lefood_artifacts(
    dataset_root: str | Path = DEFAULT_DATASET_ROOT,
    *,
    manifest_path: str | Path = DEFAULT_MANIFEST,
    source_audit_path: str | Path = DEFAULT_SOURCE_AUDIT,
    report_path: str | Path = DEFAULT_AUDIT_REPORT,
    duplicate_review_path: str | Path = DEFAULT_DUPLICATE_REVIEW,
    detect_near_duplicates: bool = True,
) -> AuditReport:
    """Create deterministic public metadata artifacts from local raw data."""

    records, source_rows, diagnostics = build_lefood_manifest(
        dataset_root, detect_near_duplicates=detect_near_duplicates
    )
    write_manifest(records, manifest_path)
    write_source_audit(source_rows, source_audit_path)
    write_duplicate_review(
        diagnostics["duplicate_candidates"], duplicate_review_path, dataset_root=dataset_root
    )
    report = audit_manifest(
        records,
        dataset_root=dataset_root,
        workbook_rows=len(source_rows),
        missing_image_rows=sum(row["status"] == "missing_image" for row in source_rows),
        exact_duplicate_components=diagnostics["exact_duplicate_components"],
        near_duplicate_components=diagnostics["near_duplicate_edges"],
        verify_hashes=True,
        manifest_sha256=sha256_file(manifest_path),
    )
    report_file = Path(report_path)
    report_file.parent.mkdir(parents=True, exist_ok=True)
    report_file.write_text(
        json.dumps(report.to_dict(), indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return report


def audit_manifest_file(
    manifest_path: str | Path,
    *,
    dataset_root: str | Path | None = None,
    verify_hashes: bool = False,
) -> AuditReport:
    """Convenience wrapper for CLI and CI."""

    return audit_manifest(
        read_manifest(manifest_path),
        dataset_root=dataset_root,
        verify_hashes=verify_hashes,
        manifest_sha256=sha256_file(manifest_path),
    )
